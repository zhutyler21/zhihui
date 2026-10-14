"""Focused backend regressions; every fixture stays in this project's work tree."""
import json
import math
from pathlib import Path
import shutil
import sqlite3
import tempfile
import threading
import time
import urllib.request
import urllib.error
from unittest.mock import patch
import unittest

ROOT = Path(__file__).resolve().parents[1]
EXPECTED = Path('/Users/zhutyler21/我的项目仓库/知会本地会议').resolve()
assert ROOT == EXPECTED, 'wrong checkout'


class FixtureModels:
    """Deterministic vectors only isolate storage/filtering, not real model quality."""
    def embed(self, texts):
        return [[float('预算' in t or '费用' in t), float('日期' in t or '上线' in t), 0.1] for t in texts]

    def status(self):
        return {'fixture':True}

    def generate(self, prompt, system, on_chunk=None, cancel_event=None, max_tokens=1600):
        self.prompts = getattr(self,'prompts',[])+[prompt]
        if '取消测试' in prompt:
            for _ in range(200):
                if cancel_event and cancel_event.is_set():
                    raise InterruptedError('取消')
                time.sleep(0.01)
        result = '合成测试回答，资料未提及其他内容。[1]'
        if on_chunk:
            on_chunk(result)
        return result

    def transcribe(self, path, cancel_event=None):
        return {'text':'合成转写预算 31472 元。','duration_seconds':4,'language':'zh'}


class StoreTests(unittest.TestCase):
    def setUp(self):
        from backend.store import Store
        self.path = Path(tempfile.mkdtemp(prefix='backend-', dir=ROOT/'work'/'fixture'))
        self.models = FixtureModels()
        self.store = Store(self.path, self.models)

    def tearDown(self):
        shutil.rmtree(self.path)

    def test_meeting_scope_excludes_other_meetings_files_and_trash(self):
        a = self.store.create_meeting('合成会议 A')
        b = self.store.create_meeting('合成会议 B')
        s = self.store.add_segment(a['id'], '预算是 31472 元，负责人李青。')
        self.store.add_segment(b['id'], '预算是 98204 元，负责人王红。')
        v = self.models.embed(['费用多少'])[0]
        found = self.store.retrieve(v, 'meeting', a['id'])
        self.assertTrue(found)
        self.assertTrue(all(x['meeting_id'] == a['id'] for x in found))
        self.assertFalse(any('98204' in x['text'] for x in found))
        self.store.segment_trash(a['id'], s['id'], True)
        self.assertEqual(self.store.retrieve(v, 'meeting', a['id']), [])
        self.store.segment_trash(a['id'], s['id'], False)
        self.assertTrue(self.store.retrieve(v, 'meeting', a['id']))
        self.store.meeting_trash(a['id'], True)
        self.assertTrue(all(x['meeting_id'] != a['id'] for x in self.store.retrieve(v, 'all')))
        with self.assertRaises(ValueError):
            self.store.retrieve(v, 'meeting', a['id'])

    def test_notes_persist_and_edits_do_not_leave_stale_index(self):
        a = self.store.create_meeting('合成持久化会议')
        self.store.patch_meeting(a['id'], {'note': '日期：11 月 9 日上线。'})
        from backend.store import Store
        reopened = Store(self.path, self.models)
        self.assertIn('11 月', reopened.get_meeting(a['id'])['note'])
        reopened.patch_meeting(a['id'], {'note': '预算：200 元。'})
        items = reopened.retrieve(self.models.embed(['费用'])[0], 'meeting', a['id'])
        self.assertFalse(any('11 月' in x['text'] for x in items))
        with self.assertRaises(ValueError):
            reopened.patch_meeting(a['id'], {'title': ''})

    def test_files_immutable_dedup_and_meetings_scope(self):
        a = self.store.create_meeting('合成附件会议')
        p = self.path/'source.txt'; p.write_text('预算 31472 元。')
        original = p.read_bytes()
        first = self.store.import_file(str(p),a['id'])
        duplicate = self.store.import_file(str(p),a['id'])
        self.assertEqual(first['id'],duplicate['id'])
        self.assertEqual(p.read_bytes(),original)
        self.assertNotEqual(Path(first['path']),p)
        p.write_text('预算 98204 元。')
        self.assertIn('31472',Path(first['path']).read_text())
        global_file = self.store.import_file(str(p))
        v = self.models.embed(['费用'])[0]
        self.assertTrue(any(x['source_id']==global_file['id'] for x in self.store.retrieve(v,'all')))
        self.assertFalse(any(x['source_id']==global_file['id'] for x in self.store.retrieve(v,'meetings')))
        self.store.source_trash(first['id'],True)
        self.assertEqual(self.store.retrieve(v,'meeting',a['id']),[])
        self.assertEqual(self.store.list_sources(a['id'],trash=True)[0]['id'],first['id'])
        self.store.source_trash(first['id'],False)
        self.assertTrue(self.store.retrieve(v,'meeting',a['id']))

    def test_file_batch_preserves_success_and_records_failure(self):
        p=self.path/'good.md'; p.write_text('日期 11 月 9 日上线。')
        result=self.store.import_files([str(self.path/'missing.txt'),str(p)])
        self.assertEqual([x['status'] for x in result['items']],['error','indexed'])
        self.assertEqual(len(self.store.list_sources()),2)

    def test_embedding_failure_preserves_text_and_reindexes(self):
        class Broken:
            def embed(self,texts): raise RuntimeError('合成模型缺失')
        self.store.models=Broken()
        a=self.store.create_meeting('合成索引失败会议')
        segment=self.store.add_segment(a['id'],'预算 31472 元。')
        self.assertIn('31472',self.store.get_meeting(a['id'])['segments'][0]['text'])
        self.assertEqual(self.store.retrieve(self.models.embed(['费用'])[0],'meeting',a['id']),[])
        self.store.models=self.models
        self.store.reindex_pending()
        self.assertTrue(self.store.retrieve(self.models.embed(['费用'])[0],'meeting',a['id']))

    def test_summary_rejects_stale_evidence(self):
        a=self.store.create_meeting('合成摘要会议')
        self.store.add_segment(a['id'],'预算 31472 元。')
        evidence=self.store.summary_evidence(a['id'])
        self.store.patch_meeting(a['id'],{'note':'日期 11 月 9 日。'})
        with self.assertRaises(ValueError):
            self.store.save_summary(a['id'],'旧摘要',evidence['content_version'])
        self.assertEqual(self.store.get_meeting(a['id'])['summary'],'')

    def test_edit_transcript_preserves_timing_and_invalidates_old_facts(self):
        a=self.store.create_meeting('合成校对会议')
        segment=self.store.add_segment(a['id'],'负责人李清，预算 11111 元。',10,20,audio_path='synthetic-original-audio')
        evidence=self.store.summary_evidence(a['id'])
        self.store.save_summary(a['id'],'旧摘要：预算 11111 元。',evidence['content_version'])
        edited=self.store.edit_segment(a['id'],segment['id'],'负责人李青，预算 31472 元。')
        self.assertEqual((edited['start_seconds'],edited['end_seconds']),(10,20))
        self.assertEqual(self.store.get_meeting(a['id'])['summary'],'')
        found=self.store.retrieve(self.models.embed(['费用'])[0],'meeting',a['id'])
        self.assertFalse(any('11111' in item['text'] for item in found))
        self.assertTrue(any('李青' in item['text'] for item in found))
        with self.store.db() as conn:
            self.assertEqual(conn.execute('SELECT audio_path FROM segments WHERE id=?',(segment['id'],)).fetchone()[0],'synthetic-original-audio')
        self.store.segment_trash(a['id'],segment['id'],True)
        with self.assertRaises(ValueError):self.store.edit_segment(a['id'],segment['id'],'不能修改回收站文字')

    def test_import_snapshot_dedup_and_source_untouched(self):
        p = self.path/'original.sqlite'
        c = sqlite3.connect(p)
        c.executescript('CREATE TABLE event_transcript_table(event_id TEXT,timestamp INTEGER,event_name TEXT,transcript TEXT,post_event_summarization TEXT); CREATE TABLE notes_table(id TEXT,title TEXT,content TEXT);')
        transcript = [{'segment_title': '合成片段', 'turns': [{'speaker':'合成甲','text':'预算 31472 元。'}], 'duration_seconds': 4}]
        c.execute('INSERT INTO event_transcript_table VALUES(?,?,?,?,?)', ('event_1000',1000,'合成导入会议',json.dumps(transcript),'合成摘要'))
        c.execute('INSERT INTO notes_table VALUES(?,?,?)', ('note_event_1000','合成笔记','日期 11 月 9 日。'))
        c.commit(); c.close()
        before = p.read_bytes()
        first = self.store.import_foresight(p)
        second = self.store.import_foresight(p)
        self.assertEqual(first['imported_meetings'], 1)
        self.assertEqual(second['imported_meetings'], 0)
        self.assertEqual(p.read_bytes(), before)
        meetings = self.store.list_meetings()
        self.assertEqual(len(meetings), 1)
        full = self.store.get_meeting(meetings[0]['id'])
        self.assertEqual(full['summary'], '合成摘要')
        self.assertIn('11 月',full['note'])
        self.assertIn('31472',full['segments'][0]['text'])

    def test_import_old_attachment_dedup_and_unknown_native_source_skip(self):
        p=self.path/'original.sqlite'; a=self.path/'old-files.sqlite'
        c=sqlite3.connect(p)
        c.executescript('CREATE TABLE event_transcript_table(event_id TEXT,timestamp INTEGER,event_name TEXT,transcript TEXT,post_event_summarization TEXT); CREATE TABLE source_text_table(document_id TEXT,document_name TEXT,content TEXT,source_link TEXT,source_type INTEGER); CREATE TABLE event_document_table(event_id TEXT,document_id TEXT);')
        c.execute('INSERT INTO event_transcript_table VALUES(?,?,?,?,?)',('event_2000',2000,'合成旧附件会议','[]','摘要'))
        c.execute('INSERT INTO source_text_table VALUES(?,?,?,?,?)',('unsafe','合成邮件','不应导入','https://mail.example.invalid',2))
        c.execute('INSERT INTO event_document_table VALUES(?,?)',('event_2000','unsafe'))
        unknown=self.path/'unknown.txt';unknown.write_text('未知类型 999 内容不能导入。')
        c.execute('INSERT INTO source_text_table VALUES(?,?,?,?,?)',('unknown','合成未知类型','不应导入',str(unknown),999))
        c.execute('INSERT INTO event_document_table VALUES(?,?)',('event_2000','unknown'))
        known=self.path/'known.txt';known.write_text('已验证类型 2 的本地文档预算 20000 元。')
        c.execute('INSERT INTO source_text_table VALUES(?,?,?,?,?)',('known','合成已知文档','原抽取文本',known.as_uri(),2))
        c.execute('INSERT INTO event_document_table VALUES(?,?)',('event_2000','known'))
        c.commit();c.close()
        f=self.path/'old.txt';f.write_text('附件预算 31472 元。')
        c=sqlite3.connect(a)
        c.execute('CREATE TABLE meeting_files(id TEXT,event_id TEXT,event_timestamp INTEGER,filename TEXT,saved_path TEXT,text TEXT)')
        c.execute('INSERT INTO meeting_files VALUES(?,?,?,?,?,?)',('fixture-file','event_2000',2000,'合成附件.txt',str(f),'附件预算 31472 元。'))
        c.commit();c.close()
        before=(p.read_bytes(),a.read_bytes(),f.read_bytes())
        result=self.store.import_foresight(p,a)
        self.assertTrue(result['warnings'])
        self.store.import_foresight(p,a)
        self.assertEqual(len(self.store.list_sources()),2)
        self.assertFalse(any('不应导入' in x['text'] or '未知类型 999' in x['text'] for x in self.store.list_sources()))
        self.assertTrue(any('已验证类型 2' in x['text'] for x in self.store.list_sources()))
        self.assertEqual(before,(p.read_bytes(),a.read_bytes(),f.read_bytes()))


class APITests(unittest.TestCase):
    def setUp(self):
        from backend.server import make_server
        self.path=Path(tempfile.mkdtemp(prefix='http-',dir=ROOT/'work'/'fixture'))
        self.models=FixtureModels()
        self.server=make_server(0,self.path/'data',self.path/'token',self.models)
        self.thread=threading.Thread(target=self.server.serve_forever,daemon=True)
        self.thread.start()
        self.base='http://127.0.0.1:'+str(self.server.server_port)

    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.server.app.close()
        self.thread.join(timeout=5)
        shutil.rmtree(self.path)

    def request(self,method,path,payload=None,authorized=True):
        headers={'Content-Type':'application/json'}
        if authorized:
            headers['Authorization']='Bearer '+self.server.token
        req=urllib.request.Request(self.base+path,data=json.dumps(payload).encode() if payload is not None else None,headers=headers,method=method)
        try:
            with urllib.request.urlopen(req,timeout=10) as result:
                return result.status,json.load(result)
        except urllib.error.HTTPError as exc:
            return exc.code,json.load(exc)

    def wait_job(self,job_id):
        for _ in range(100):
            status,job=self.request('GET','/api/jobs/'+job_id)
            if job['state'] in ('done','error','cancelled'):
                return job
            time.sleep(0.03)
        self.fail('job did not finish')

    def test_authenticated_entry_point_and_validation(self):
        self.assertEqual(self.request('GET','/api/health',authorized=False)[0],401)
        status,health=self.request('GET','/api/health')
        self.assertEqual(status,200);self.assertEqual(health['app'],'zhihui');self.assertEqual(health['version'],1)
        self.assertEqual((self.path/'token').stat().st_mode&0o777,0o600)
        self.assertEqual(self.request('POST','/api/meetings',{'title':''})[0],400)
        self.assertEqual(self.request('GET','/api/unknown')[0],404)
        self.assertEqual(self.request('POST','/api/ask',{'scope':'meeting','question':'费用？'})[0],400)
        self.assertEqual(self.request('POST','/api/sources',{'paths':['relative.txt']})[0],400)
        self.assertEqual(self.request('POST','/api/ask',{'scope':'all','question':'预算','history':['bad']})[0],400)

    def test_empty_evidence_does_not_use_history_or_generate(self):
        status,result=self.request('POST','/api/ask',{'question':'多少钱？','scope':'all','history':[{'role':'assistant','content':'预算 99999 元。'}]})
        self.assertEqual(status,200)
        job=self.wait_job(result['job_id'])
        self.assertEqual(job['state'],'done')
        self.assertEqual(job['result']['citations'],[])
        self.assertNotIn('99999',job['result']['answer'])
        self.assertFalse(hasattr(self.models,'prompts'))

    def test_scoped_answer_native_jobs_and_cancel(self):
        _,a=self.request('POST','/api/meetings',{'title':'合成 HTTP 会议 A'})
        _,b=self.request('POST','/api/meetings',{'title':'合成 HTTP 会议 B'})
        self.request('POST','/api/meetings/'+a['id']+'/segments',{'text':'预算 31472 元。'})
        self.request('POST','/api/meetings/'+b['id']+'/segments',{'text':'预算 98204 元。'})
        _,result=self.request('POST','/api/ask',{'question':'费用多少？','scope':'meeting','meeting_id':a['id'],'history':[{'role':'assistant','content':'预算 98204 元。'}]})
        job=self.wait_job(result['job_id']);self.assertEqual(job['state'],'done')
        self.assertTrue(all(x['meeting_id']==a['id'] for x in job['result']['citations']))
        self.assertTrue(all(not key.startswith('_') for item in job['result']['citations'] for key in item))
        self.assertNotIn('98204',self.models.prompts[-1])
        _,result=self.request('POST','/api/ask',{'question':'取消测试费用','scope':'meeting','meeting_id':a['id']})
        for _ in range(100):
            _,current=self.request('GET','/api/jobs/'+result['job_id'])
            if current['state']=='running':break
            time.sleep(0.01)
        self.request('POST','/api/jobs/'+result['job_id']+'/cancel',{})
        self.assertEqual(self.wait_job(result['job_id'])['state'],'cancelled')

    def test_transcribe_idempotent_request_and_hash_binding(self):
        _,a=self.request('POST','/api/meetings',{'title':'合成录音会议'})
        p=self.path/'fixture.wav';p.write_bytes(b'fixture-asr-not-real-audio')
        request={'path':str(p),'request_id':'synthetic-request-1','start_seconds':10}
        endpoint='/api/meetings/'+a['id']+'/transcribe'
        _,first=self.request('POST',endpoint,request)
        job=self.wait_job(first['job_id']);self.assertEqual(job['state'],'done')
        self.assertEqual(job['result']['duration_seconds'],4)
        _,second=self.request('POST',endpoint,request)
        self.assertEqual(first['job_id'],second['job_id'])
        _,meeting=self.request('GET','/api/meetings/'+a['id'])
        self.assertEqual(len(meeting['segments']),1)
        self.assertEqual(meeting['segments'][0]['start_seconds'],10)
        p.write_bytes(b'changed-content')
        self.assertEqual(self.request('POST',endpoint,request)[0],400)

    def test_long_summary_processes_all_evidence(self):
        _,a=self.request('POST','/api/meetings',{'title':'合成长会议'})
        text='预算事实。'*2500+'最后行动项：李青在 11 月 9 日上线。'
        self.request('POST','/api/meetings/'+a['id']+'/segments',{'text':text})
        _,result=self.request('POST','/api/meetings/'+a['id']+'/summarize',{})
        job=self.wait_job(result['job_id']);self.assertEqual(job['state'],'done')
        self.assertGreater(job['result']['evidence_groups'],1)
        self.assertFalse(job['result']['sampled'])
        self.assertTrue(any('最后行动项' in p for p in self.models.prompts))
        self.assertTrue(all(len(p)<13_000 for p in self.models.prompts))

    def test_ask_generation_rejects_edit_or_trash_and_clears_partial(self):
        for mutation in ('edit','segment-trash','meeting-trash','meeting-roundtrip'):
            with self.subTest(mutation=mutation):
                _,a=self.request('POST','/api/meetings',{'title':'合成并发会议 '+mutation})
                _,segment=self.request('POST','/api/meetings/'+a['id']+'/segments',{'text':'预算 11111 元。'})
                entered,released=threading.Event(),threading.Event()
                real=self.models.generate
                def blocked(prompt,system,on_chunk=None,cancel_event=None,max_tokens=1600):
                    if on_chunk:on_chunk('旧预算 11111 元。[1]')
                    entered.set()
                    if not released.wait(5):raise RuntimeError('fixture barrier timeout')
                    return '旧预算 11111 元。[1]'
                self.models.generate=blocked
                try:
                    _,submitted=self.request('POST','/api/ask',{'question':'预算？','scope':'meeting','meeting_id':a['id']})
                    self.assertTrue(entered.wait(3))
                    if mutation=='edit':
                        status,_=self.request('PATCH','/api/meetings/'+a['id']+'/segments/'+segment['id'],{'text':'已更正预算 22222 元。'})
                    elif mutation=='segment-trash':
                        status,_=self.request('POST','/api/meetings/'+a['id']+'/segments/'+segment['id']+'/trash',{})
                    else:
                        status,_=self.request('POST','/api/meetings/'+a['id']+'/trash',{})
                        if mutation=='meeting-roundtrip':
                            self.assertEqual(self.request('POST','/api/meetings/'+a['id']+'/restore',{})[0],200)
                    self.assertEqual(status,200)
                    released.set();job=self.wait_job(submitted['job_id'])
                    self.assertEqual(job['state'],'error')
                    self.assertIsNone(job['result'])
                    self.assertEqual(job['partial'],'')
                    self.assertIn('资料已变化',job['error'])
                finally:
                    released.set();self.models.generate=real

    def test_cancel_after_transcript_or_summary_commit_reports_saved_result(self):
        for kind in ('transcribe','summarize'):
            with self.subTest(kind=kind):
                _,a=self.request('POST','/api/meetings',{'title':'合成提交边界 '+kind})
                if kind=='summarize':
                    self.request('POST','/api/meetings/'+a['id']+'/segments',{'text':'预算 31472 元。'})
                entered,released=threading.Event(),threading.Event()
                real=self.models.embed
                marker='合成转写' if kind=='transcribe' else '合成测试回答'
                def blocked(texts):
                    if any(marker in text for text in texts):
                        entered.set()
                        if not released.wait(5):raise RuntimeError('fixture barrier timeout')
                    return real(texts)
                self.models.embed=blocked
                try:
                    if kind=='transcribe':
                        p=self.path/(kind+'.wav');p.write_bytes(b'fixture-boundary-asr')
                        payload={'path':str(p),'request_id':'boundary'}
                    else:payload={}
                    _,submitted=self.request('POST','/api/meetings/'+a['id']+'/'+kind,payload)
                    self.assertTrue(entered.wait(3))
                    self.request('POST','/api/jobs/'+submitted['job_id']+'/cancel',{})
                    released.set();job=self.wait_job(submitted['job_id'])
                    self.assertEqual(job['state'],'done')
                    self.assertEqual(job['partial'],'')
                    self.assertTrue(job['result']['warnings'])
                    self.assertIn('已保存',job['message'])
                    _,current=self.request('GET','/api/meetings/'+a['id'])
                    if kind=='transcribe':
                        self.assertIsNotNone(job['result']['segment'])
                        self.assertEqual(len(current['segments']),1)
                    else:
                        self.assertTrue(current['summary'])
                        self.assertEqual(job['result']['summary'],current['summary'])
                finally:released.set();self.models.embed=real

    def test_file_copy_cleans_unpublished_copy_after_meeting_trash(self):
        from backend import store as store_module
        _,a=self.request('POST','/api/meetings',{'title':'合成复制并发会议'})
        p=self.path/'original.txt';p.write_text('合成原文件预算 31472 元。')
        original=p.read_bytes();entered,released=threading.Event(),threading.Event()
        real=store_module.extract_text
        def blocked(path):
            entered.set()
            if not released.wait(5):raise RuntimeError('fixture barrier timeout')
            return real(path)
        with patch.object(store_module,'extract_text',blocked):
            _,submitted=self.request('POST','/api/sources',{'paths':[str(p)],'meeting_id':a['id']})
            try:
                self.assertTrue(entered.wait(3))
                self.request('POST','/api/meetings/'+a['id']+'/trash',{})
                released.set();job=self.wait_job(submitted['job_id'])
                self.assertEqual(job['state'],'error')
                self.assertEqual(p.read_bytes(),original)
                self.assertEqual(list((self.path/'data/attachments').glob('*')),[])
                with self.server.app.store.db() as conn:
                    self.assertEqual(conn.execute("SELECT COUNT(*) FROM sources WHERE kind='file'").fetchone()[0],0)
            finally:released.set()

    def test_silent_transcribe_is_done_and_retains_audio_without_segment(self):
        _,a=self.request('POST','/api/meetings',{'title':'合成静音会议'})
        p=self.path/'silent.wav';p.write_bytes(b'fixture-silence-not-real-audio')
        self.models.transcribe=lambda path,cancel_event=None:{'text':'','duration_seconds':10,'language':None}
        _,submitted=self.request('POST','/api/meetings/'+a['id']+'/transcribe',{'path':str(p),'request_id':'silent'})
        job=self.wait_job(submitted['job_id']);self.assertEqual(job['state'],'done')
        self.assertIsNone(job['result']['segment']);self.assertEqual(job['result']['duration_seconds'],10)
        self.assertEqual(len(list((self.path/'data/recordings').glob('*'))),1)
        self.assertEqual(self.server.app.store.get_meeting(a['id'])['segments'],[])

    def test_cancel_before_file_commit_cleans_copy_and_original_stays(self):
        from backend import store as store_module
        p=self.path/'cancelled-file.txt';p.write_text('合成取消文件')
        entered,released=threading.Event(),threading.Event();real=store_module.extract_text
        def blocked(path):
            entered.set()
            if not released.wait(5):raise RuntimeError('fixture barrier timeout')
            return real(path)
        with patch.object(store_module,'extract_text',blocked):
            _,submitted=self.request('POST','/api/sources',{'paths':[str(p)]})
            try:
                self.assertTrue(entered.wait(3))
                self.request('POST','/api/jobs/'+submitted['job_id']+'/cancel',{})
                released.set();job=self.wait_job(submitted['job_id'])
                self.assertEqual(job['state'],'cancelled');self.assertEqual(job['partial'],'')
                self.assertEqual(p.read_text(),'合成取消文件')
                self.assertEqual(list((self.path/'data/attachments').glob('*')),[])
                self.assertEqual(self.server.app.store.list_sources(),[])
            finally:released.set()

    def test_stale_summary_error_clears_streamed_draft(self):
        _,a=self.request('POST','/api/meetings',{'title':'合成摘要草稿并发'})
        self.request('POST','/api/meetings/'+a['id']+'/segments',{'text':'预算 11111 元。'})
        entered,released=threading.Event(),threading.Event()
        def blocked(prompt,system,on_chunk=None,cancel_event=None,max_tokens=1600):
            if on_chunk:on_chunk('旧摘要预算 11111 元。')
            entered.set()
            if not released.wait(5):raise RuntimeError('fixture barrier timeout')
            return '旧摘要预算 11111 元。'
        real=self.models.generate;self.models.generate=blocked
        try:
            _,submitted=self.request('POST','/api/meetings/'+a['id']+'/summarize',{})
            self.assertTrue(entered.wait(3))
            self.request('PATCH','/api/meetings/'+a['id'],{'note':'最新预算 22222 元。'})
            released.set();job=self.wait_job(submitted['job_id'])
            self.assertEqual(job['state'],'error');self.assertEqual(job['partial'],'')
            self.assertIsNone(job['result'])
            self.assertEqual(self.server.app.store.get_meeting(a['id'])['summary'],'')
        finally:released.set();self.models.generate=real


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.path=Path(tempfile.mkdtemp(prefix='documents-',dir=ROOT/'work'/'fixture'))

    def tearDown(self):
        shutil.rmtree(self.path)

    def test_real_office_extraction_and_pdf_ocr_status(self):
        from backend.documents import extract_text
        from docx import Document
        from openpyxl import Workbook
        from pypdf import PdfWriter
        import zipfile
        doc=Document();doc.add_paragraph('合成会议预算 31472 元。');doc.add_table(rows=1,cols=1).cell(0,0).text='负责人李青'
        doc.save(self.path/'fixture.docx')
        text,_=extract_text(self.path/'fixture.docx');self.assertIn('31472',text);self.assertIn('李青',text)
        workbook=Workbook();workbook.active.append(['日期','11 月 9 日']);workbook.save(self.path/'fixture.xlsx')
        self.assertIn('11 月',extract_text(self.path/'fixture.xlsx')[0])
        with zipfile.ZipFile(self.path/'fixture.pptx','w') as archive:
            archive.writestr('ppt/slides/slide1.xml','<p:sld xmlns:p="p" xmlns:a="a"><a:t>合成上线日期</a:t></p:sld>')
        self.assertIn('上线日期',extract_text(self.path/'fixture.pptx')[0])
        writer=PdfWriter();writer.add_blank_page(width=100,height=100);writer.write(self.path/'blank.pdf')
        text,note=extract_text(self.path/'blank.pdf');self.assertEqual(text,'');self.assertIn('OCR',note)
        (self.path/'fixture.html').write_text('<script>不可见内容</script><p>会议预算</p>')
        text,_=extract_text(self.path/'fixture.html');self.assertIn('会议预算',text);self.assertNotIn('不可见内容',text)

    def test_corrupt_and_unsupported_inputs(self):
        from backend.documents import extract_text
        (self.path/'bad.pdf').write_bytes(b'not-a-pdf')
        with self.assertRaises(Exception):extract_text(self.path/'bad.pdf')
        (self.path/'unknown.bin').write_bytes(b'foo')
        self.assertEqual(extract_text(self.path/'unknown.bin')[0],'')


if __name__ == '__main__':
    unittest.main()
