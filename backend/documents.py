"""Local file extraction. No network, original files are never modified."""
from html.parser import HTMLParser
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET

MAX_FILE_BYTES = 256 * 1024 * 1024
MAX_TEXT_CHARS = 500_000
MAX_ZIP_BYTES = 80 * 1024 * 1024


class PlainHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.ignore = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'):
            self.ignore += 1
        elif tag in ('p', 'div', 'br', 'li', 'tr', 'h1', 'h2', 'h3'):
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in ('script', 'style'):
            self.ignore = max(0, self.ignore - 1)

    def handle_data(self, data):
        if not self.ignore:
            self.parts.append(data)


def _decoded(path):
    data = path.read_bytes()
    encodings = ('utf-16',) if data.startswith((b'\xff\xfe', b'\xfe\xff')) else ('utf-8-sig', 'gb18030')
    for encoding in encodings:
        try:
            text = data.decode(encoding)
            if '\x00' in text:
                raise ValueError('文件包含二进制内容，无法作为文字读取。')
            return text
        except UnicodeDecodeError:
            pass
    raise ValueError('无法识别文件文字编码，请另存为 UTF-8。')


def _zip_check(path):
    with zipfile.ZipFile(path) as archive:
        if sum(x.file_size for x in archive.infolist()) > MAX_ZIP_BYTES:
            raise ValueError('Office 文件解压后的内容超过 80 MB。')
        if any(x.file_size > 20 * 1024 * 1024 for x in archive.infolist()):
            raise ValueError('Office 文件中的单项内容过大。')


def extract_text(path):
    """Return (text, Chinese read-status). Empty text is a visible unreadable file."""
    path = Path(path)
    if not path.is_file():
        raise ValueError('文件不存在或不是普通文件。')
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError('文件超过 256 MB，请拆分后导入。')
    suffix = path.suffix.lower()
    text = ''
    if suffix in ('.txt', '.md', '.csv', '.json', '.log', '.html', '.htm'):
        text = _decoded(path)
        if suffix in ('.html', '.htm'):
            parser = PlainHTML()
            parser.feed(text)
            text = ''.join(parser.parts)
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('PDF 有密码，请解密后导入。')
        if len(reader.pages) > 500:
            raise ValueError('PDF 超过 500 页，请拆分后导入。')
        parts, total = [], 0
        for page in reader.pages:
            part = page.extract_text() or ''
            parts.append(part)
            total += len(part)
            if total > MAX_TEXT_CHARS:
                break
        text = '\n\n'.join(parts)
        if not text.strip():
            return '', 'PDF 未读取到文字，可能为扫描件，需要先进行 OCR。'
    elif suffix == '.docx':
        _zip_check(path)
        from docx import Document
        document = Document(str(path))
        # Native block order preserves paragraphs interspersed with tables.
        parts = []
        for child in document.element.body:
            if child.tag.endswith('}p'):
                parts.append(''.join(x.text or '' for x in child.iter() if x.tag.endswith('}t')))
            elif child.tag.endswith('}tbl'):
                for row in child:
                    if row.tag.endswith('}tr'):
                        parts.append('\t'.join(''.join(x.text or '' for x in cell.iter() if x.tag.endswith('}t')) for cell in row if cell.tag.endswith('}tc')))
        text = '\n'.join(parts)
    elif suffix == '.pptx':
        _zip_check(path)
        parts = []
        with zipfile.ZipFile(path) as archive:
            slides = sorted((n for n in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)), key=lambda n: int(re.search(r'slide(\d+)', n).group(1)))
            for index, name in enumerate(slides, 1):
                root = ET.fromstring(archive.read(name))
                parts.append(f'第 {index} 页\n' + '\n'.join(x.text or '' for x in root.iter() if x.tag.endswith('}t')))
        text = '\n\n'.join(parts)
    elif suffix == '.xlsx':
        _zip_check(path)
        from openpyxl import load_workbook
        workbook = load_workbook(str(path), read_only=True, data_only=False, keep_links=False)
        try:
            parts, total = [], 0
            for sheet in workbook:
                parts.append(f'工作表：{sheet.title}')
                for row in sheet.iter_rows(values_only=True):
                    part = '\t'.join('' if value is None else str(value) for value in row)
                    if part.strip():
                        parts.append(part)
                        total += len(part)
                    if total > MAX_TEXT_CHARS:
                        break
                if total > MAX_TEXT_CHARS:
                    break
            text = '\n'.join(parts)
        finally:
            workbook.close()
    else:
        return '', '文件副本已保存，此类型暂不支持文字读取。'
    text = text.replace('\x00', '').strip()
    if not text:
        return '', '文件副本已保存，但未读取到文字。'
    if len(text) > MAX_TEXT_CHARS:
        return text[:MAX_TEXT_CHARS], '文字较多，已读取前 50 万字。'
    return text, '文字已读取。'


def chunks(text, size=420, overlap=60):
    """Bounded overlapping evidence units, split near punctuation where practical."""
    result = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundaries = [text.rfind(mark, start + size//2, end) for mark in ('\n', '。', '！', '？', '. ')]
            split = max(boundaries)
            if split > start:
                end = split + 1
        part = text[start:end].strip()
        if part:
            result.append(part)
        if end == len(text):
            break
        start = max(start + 1, end - overlap)
    return result
